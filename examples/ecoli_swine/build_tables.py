"""Derive the E. coli input tables from Supplementary Data 1 and 13 of Li et al.

Li X, Hu H, Zhu Y, et al. Population structure and antibiotic resistance of
swine extraintestinal pathogenic Escherichia coli from China. Nature
Communications 2024;15:5811. doi:10.1038/s41467-024-50268-2 (CC BY 4.0).
The workbook of Supplementary Data 1-14 is the file 41467_2024_50268_MOESM4_ESM.xlsx
served beside the article; its digest is pinned below.

    python examples/ecoli_swine/build_tables.py <path to the workbook>

Six agents are kept, each read on the broth microdilution range the source
reports (BD Phoenix; the "Range" block of Supplementary Data 13). A reading
written "<=1" is at or below the lowest well, ">16" above the highest, and an
unprefixed value is the dilution interval ending at that concentration. Four
isolates are set aside before any analysis: A24 (no gentamicin reading), A62
and A71 (no ciprofloxacin reading) and A154 (meropenem 3 mg/L, which lies on
no well of the range), so that every agent is analysed on the same isolates.
Isolates whose MLST is "-" in Supplementary Data 1 keep a blank lineage and
are counted by the run as untyped.

Three reading tables are written: the full panel, the same readings merged
into two categories at one prespecified well per agent, and the full panel
without the isolates of ST410, the most frequent sequence type.
"""
from __future__ import annotations

import csv
import hashlib
import json
import sys
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
DATA = HERE / "data"
WORKBOOK_SHA256 = "b723215ec2dd20d889a7284858620094bd39b960f4ff5cdca9bee3365a5d27c9"
PANEL = "BD Phoenix"
AGENTS = {
    "Cefotaxime": ("cefotaxime", [1, 2, 4, 8, 16, 32], 4),
    "Ceftazidime": ("ceftazidime", [1, 2, 4, 8, 16], 4),
    "Meropenem": ("meropenem", [1, 2, 4, 8], 2),
    "Ciprofloxacin": ("ciprofloxacin", [0.5, 1, 2], 1),
    "Gentamicin": ("gentamicin", [2, 4, 8], 4),
    "Tetracycline": ("tetracycline", [2, 4, 8], 4),
}
SET_ASIDE = {"A24": "no gentamicin reading", "A62": "no ciprofloxacin reading",
             "A71": "no ciprofloxacin reading", "A154": "meropenem 3 mg/L, on no well of the range"}
EXCLUDED_ST = "410"


def parse(text: str):
    """A reading as (operator, value): '<=1' -> ('<=', 1.0), '>16' -> ('>', 16.0), '8' -> ('', 8.0)."""
    text = str(text).strip().replace("≤", "<=").replace("≥", ">=")
    for op in ("<=", ">=", "<", ">"):
        if text.startswith(op):
            return op, float(text[len(op):])
    return "", float(text)


def fmt(value: float) -> str:
    return str(int(value)) if float(value).is_integer() else str(value)


def main(workbook: Path) -> int:
    digest = hashlib.sha256(workbook.read_bytes()).hexdigest()
    if digest != WORKBOOK_SHA256:
        raise SystemExit(f"workbook digest {digest} is not the pinned {WORKBOOK_SHA256}")
    wb = openpyxl.load_workbook(workbook, read_only=True)
    rows1 = list(wb["Supplementary Data 1"].iter_rows(values_only=True))
    rows13 = list(wb["Supplementary Data 13"].iter_rows(values_only=True))
    head1, head13 = rows1[1], rows13[1]
    col1 = {name: i for i, name in enumerate(head1) if name}
    col13 = {name: i for i, name in enumerate(head13) if name}

    metadata = []
    for r in rows1[2:]:
        if not r[col1["Strain name"]]:
            continue
        st = str(r[col1["MLST"]]).strip()
        metadata.append(dict(
            isolate_id=str(r[col1["Strain name"]]).strip(),
            st="" if st == "-" else f"ST{st}",
            phylogroup=str(r[col1["Phylogroup"]]).strip(),
            isolation_year=str(r[col1["Isolation time (Year)"]]).strip(),
            tissue=str(r[col1["Tissue source"]]).strip(),
            province=str(r[col1["Province in china"]]).strip(),
            serotype=str(r[col1["Combination of H- and O- serotype"]]).strip(),
            biosample=str(r[col1["BioSample accession number"]]).strip(),
        ))
    DATA.mkdir(exist_ok=True)
    with (DATA / "metadata.csv").open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(metadata[0]))
        w.writeheader()
        w.writerows(metadata)

    full, two = [], []
    counts = dict(ast_rows=0, set_aside={}, readings={})
    for r in rows13[2:]:
        isolate = r[col13["Strain"]]
        if not isolate:
            continue
        isolate = str(isolate).strip()
        counts["ast_rows"] += 1
        if isolate in SET_ASIDE:
            counts["set_aside"][isolate] = SET_ASIDE[isolate]
            continue
        for source_name, (agent, wells, cut) in AGENTS.items():
            op, value = parse(r[col13[source_name]])
            full.append(dict(isolate_id=isolate, agent=agent, operator=op, measurement=fmt(value),
                             unit="mg/L", panel=PANEL))
            above = (op == ">" and value >= cut) or (op == "" and value > cut) or (op == ">=" and value > cut)
            two.append(dict(isolate_id=isolate, agent=agent, operator=">" if above else "<=",
                            measurement=fmt(cut), unit="mg/L", panel=PANEL))
            counts["readings"][agent] = counts["readings"].get(agent, 0) + 1
    in_st410 = {m["isolate_id"] for m in metadata if m["st"] == f"ST{EXCLUDED_ST}"}
    without = [row for row in full if row["isolate_id"] not in in_st410]
    for name, table in (("mic_long.csv", full), ("mic_two_categories.csv", two),
                        ("mic_long_without_st410.csv", without)):
        with (DATA / name).open("w", newline="", encoding="utf-8") as f:
            w = csv.DictWriter(f, fieldnames=list(table[0]))
            w.writeheader()
            w.writerows(table)
    receipt = dict(workbook=workbook.name, workbook_sha256=digest, isolates_in_data_1=len(metadata),
                   ast_rows_in_data_13=counts["ast_rows"], set_aside=counts["set_aside"],
                   readings_per_agent=counts["readings"],
                   wells={a: w for _, (a, w, _) in AGENTS.items()},
                   two_category_cut={a: c for _, (a, _, c) in AGENTS.items()},
                   excluded_sequence_type=f"ST{EXCLUDED_ST}", isolates_of_excluded_type=len(in_st410))
    (DATA / "build_receipt.json").write_text(json.dumps(receipt, indent=1) + "\n")
    print(json.dumps(receipt, indent=1))
    return 0


if __name__ == "__main__":
    sys.exit(main(Path(sys.argv[1])))
